"""
make_table1_structural.py

Table 1 (2 mm skin depth) — Structural assessment at k=7, 2 scenarios × 2 methods.
Only k=7 available for v2; the 6-condition table from v1 is reduced to 2 conditions.

Outputs:
  publication/figures/table1_structural.xlsx
  publication/figures/table1_structural.tex

Run from project root:
  python3 publication/scripts/make_table1_structural.py
"""

import os
import numpy as np
import pandas as pd

ROOT     = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..")
CALIB    = os.path.join(ROOT, "calibration")
OUT_DIR  = os.path.join(ROOT, "publication", "figures")
XLSX_OUT = os.path.join(OUT_DIR, "table1_structural.xlsx")
TEX_OUT  = os.path.join(OUT_DIR, "table1_structural.tex")

os.makedirs(OUT_DIR, exist_ok=True)

# ── Load v2 k=7 CSVs ─────────────────────────────────────────────────────────

CSV_MAP = {
    ("S2", 7): (os.path.join(CALIB, "comparison_structure_metrics_v2_s2ms.csv"),
                "CMA-ES-S2MS", "TabICL-MS-v2"),
    ("S1", 7): (os.path.join(CALIB, "comparison_structure_metrics_v2_s1ms.csv"),
                "CMA-ES-S1MS", "TabICL-MS-v2"),
}

METRICS = ["ss_rld", "ss_angle"]

rows = []
for (scen, k), (path, cmae_label, tab_label) in CSV_MAP.items():
    df = pd.read_csv(path)
    for method_key, label in [("CMA-ES", cmae_label), ("TabICL-MS", tab_label)]:
        sub = df[df["method"] == label]
        entry = {"scenario": scen, "k": k, "method": method_key}
        for met in METRICS:
            vals = sub[met].dropna()
            entry[f"{met}_mean"] = vals.mean()
            entry[f"{met}_sd"]   = vals.std()
            entry[f"{met}_n"]    = len(vals)
        rows.append(entry)

results_df = pd.DataFrame(rows)

# ── Determine bold (best per scenario×k) ─────────────────────────────────────

results_df["bold_ss_rld"]   = False
results_df["bold_ss_angle"] = False

for (scen, k), grp in results_df.groupby(["scenario", "k"]):
    results_df.loc[grp["ss_rld_mean"].idxmax(),   "bold_ss_rld"]   = True
    results_df.loc[grp["ss_angle_mean"].idxmax(), "bold_ss_angle"] = True

# ── LaTeX table ───────────────────────────────────────────────────────────────

def fmt_mean_sd(mean, sd, bold=False, fmt=".3f"):
    s = f"{mean:{fmt}} ± {sd:{fmt}}"
    if bold:
        s = r"\textbf{" + s + "}"
    return s

tex_lines = [
    r"\begin{table}[htbp]",
    r"\centering",
    r"\caption{Structural recovery at $k$=7 with 0.2 cm skin depth (N=20 test cases for S2, N=20 for S1). "
    r"SS$_\mathrm{RLD}$: skill score for root length density; SS$_\mathrm{angle}$: skill score for crossing "
    r"angle. Skill scores $>0$ indicate better than the null predictor. Bold: best per scenario. Values: mean $\pm$ SD.}",
    r"\label{tab:structural_recovery_v2}",
    r"\begin{tabular}{llrr}",
    r"\toprule",
    r"Scenario & Method & SS$_\mathrm{RLD}$ & SS$_\mathrm{angle}$ \\",
    r"\midrule",
]

order = [("S2", 7), ("S1", 7)]
prev_scen = None

for i, (scen, k) in enumerate(order):
    sub = results_df[(results_df["scenario"] == scen) & (results_df["k"] == k)]
    if prev_scen is not None and scen != prev_scen:
        tex_lines.append(r"\midrule")
    prev_scen = scen

    scen_label = r"S2 (elong.\ known)" if scen == "S2" else r"S1 (all unknown)"
    first_row = True

    for _, row in sub.iterrows():
        method_label = "CMA-ES" if row["method"] == "CMA-ES" else "TabICL-MS"
        rld_str = fmt_mean_sd(row["ss_rld_mean"],   row["ss_rld_sd"],   bold=row["bold_ss_rld"])
        ang_str = fmt_mean_sd(row["ss_angle_mean"], row["ss_angle_sd"], bold=row["bold_ss_angle"])

        if first_row:
            scen_col = scen_label
            first_row = False
        else:
            scen_col = ""

        tex_lines.append(f"  {scen_col} & {method_label} & {rld_str} & {ang_str} \\\\")

tex_lines += [
    r"\bottomrule",
    r"\end{tabular}",
    r"\end{table}",
]

with open(TEX_OUT, "w") as f:
    f.write("\n".join(tex_lines) + "\n")
print(f"Saved LaTeX → {TEX_OUT}")

# Print preview
for (scen, k), grp in results_df.groupby(["scenario", "k"]):
    for _, row in grp.iterrows():
        print(f"{scen} k={k} {row['method']:<10}: "
              f"SS_rld={row['ss_rld_mean']:+.3f}±{row['ss_rld_sd']:.3f}  "
              f"SS_angle={row['ss_angle_mean']:+.3f}±{row['ss_angle_sd']:.3f}  "
              f"N={int(row['ss_rld_n'])}")

# ── Excel table ───────────────────────────────────────────────────────────────

SCEN_DISPLAY = {"S2": "S2 — elongation known (k=7, 0.2 cm)",
                "S1": "S1 — all unknown (k=7, 0.2 cm)"}

excel_rows = []
for (scen, k) in order:
    sub = results_df[(results_df["scenario"] == scen) & (results_df["k"] == k)]
    first_row = True
    for _, row in sub.iterrows():
        method_label = "CMA-ES" if row["method"] == "CMA-ES" else "TabICL-MS"
        excel_rows.append({
            "Scenario":        SCEN_DISPLAY[scen] if first_row else "",
            "Method":          method_label,
            "SS_RLD (mean)":   round(row["ss_rld_mean"],   3),
            "SS_RLD (SD)":     round(row["ss_rld_sd"],     3),
            "SS_angle (mean)": round(row["ss_angle_mean"], 3),
            "SS_angle (SD)":   round(row["ss_angle_sd"],   3),
            "N":               int(row["ss_rld_n"]),
            "Best SS_RLD":     row["bold_ss_rld"],
            "Best SS_angle":   row["bold_ss_angle"],
        })
        first_row = False

excel_df = pd.DataFrame(excel_rows)

with pd.ExcelWriter(XLSX_OUT, engine="openpyxl") as writer:
    excel_df.to_excel(writer, index=False, sheet_name="Table1_v2")
    ws = writer.sheets["Table1_v2"]

    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    HEADER_FILL = PatternFill("solid", fgColor="2C3E50")
    S2_FILL     = PatternFill("solid", fgColor="DCE8F5")
    S1_FILL     = PatternFill("solid", fgColor="FDE8D8")
    BOLD_FONT   = Font(bold=True)
    HEADER_FONT = Font(bold=True, color="FFFFFF")
    THIN_BORDER = Border(
        bottom=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
    )

    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    col_widths = [30, 12, 14, 12, 16, 12, 6, 12, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for row_idx, row_data in enumerate(excel_rows, start=2):
        actual_scen = "S2" if row_idx <= 3 else "S1"
        for col_idx, col_name in enumerate(excel_df.columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER
            if col_idx <= 2:
                cell.fill = S2_FILL if actual_scen == "S2" else S1_FILL
            is_bold = ((col_name == "SS_RLD (mean)"   and row_data["Best SS_RLD"]) or
                       (col_name == "SS_angle (mean)" and row_data["Best SS_angle"]))
            if is_bold:
                cell.font = BOLD_FONT

    for col_idx in [8, 9]:
        ws.column_dimensions[get_column_letter(col_idx)].hidden = True

    ws.freeze_panes = "A2"

    note_row = len(excel_rows) + 3
    ws.cell(row=note_row, column=1,
            value="Notes: skin_depth = 0.2 cm (2 mm), k = 7 seeds. "
                  "SS > 0 = better than null predictor. Bold = best per scenario. "
                  "S2: elongation known, 4 arch params. S1: all 8 params unknown.")
    ws.cell(row=note_row, column=1).font = Font(italic=True, size=9)
    ws.merge_cells(start_row=note_row, start_column=1,
                   end_row=note_row, end_column=7)

print(f"Saved Excel → {XLSX_OUT}")
